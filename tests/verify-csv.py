"""Independent XLSX verification. Run after node tests/csv-excel.test.mjs."""
import json
from pathlib import Path
import openpyxl

fixtures = Path(__file__).resolve().parent.parent / 'work' / 'csv-validation'
expected = json.loads((fixtures / 'expected.json').read_text(encoding='utf-8'))
book = openpyxl.load_workbook(fixtures / 'edge-cases.xlsx')
sheet = book['Data']
assert sheet.max_row == len(expected)
assert sheet.max_column == len(expected[0])
for r, row in enumerate(expected, 1):
    for c, value in enumerate(row, 1):
        cell = sheet.cell(r, c)
        assert cell.value == value, (r, c, repr(value), repr(cell.value))
        assert cell.data_type == 's', (r, c, cell.data_type)
assert book.sheetnames == ['Data', 'Conversion notes']
print(f'Independent openpyxl readback passed: {sheet.max_row * sheet.max_column} text cells.')
